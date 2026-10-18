"""Local-only job feedback, application tracking and resume/document helpers.

Personal resumes are stored exclusively in the Railway SQLite volume, not in Git.
No network calls happen here. Database access must be made for the allowed chat.
"""
import io
import json
import re
import time
from dataclasses import asdict
from html import escape

from docx import Document
from docx.shared import Pt, Inches
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from sources import Job

ALLOWED_FEEDBACK = frozenset(("like", "dislike"))
APPLICATION_STATUSES = {
    "saved": "📌 Хочу откликнуться",
    "applied": "✅ Откликнулся",
    "hr": "📞 Общение с HR",
    "interview": "🧠 Интервью",
    "offer": "🎉 Оффер",
    "rejected": "❌ Отказ",
}
TOPICS = {
    "game": r"игров|геймдев|gaming|game|gameplay",
    "product": r"продукт|product",
    "business": r"бизнес|business|commercial|коммерч",
    "data": r"data|данных|\bbi\b|intelligence",
    "growth": r"growth|revenue|монетиз",
    "marketing": r"маркетинг|marketing|performance",
    "finance": r"финанс|financial|margin|profit",
    "operations": r"операцион|operations|логист|supply",
    "ecommerce": r"e-commerce|ecommerce|маркетплейс|marketplace|retail",
}


def init_tables(db):
    db.execute("""CREATE TABLE IF NOT EXISTS job_snapshots (
        job_key TEXT PRIMARY KEY, payload TEXT NOT NULL, updated REAL NOT NULL)""")
    db.execute("""CREATE TABLE IF NOT EXISTS feedback (
        chat_id INTEGER NOT NULL, job_key TEXT NOT NULL, vote TEXT NOT NULL,
        updated REAL NOT NULL, PRIMARY KEY(chat_id, job_key))""")
    db.execute("""CREATE TABLE IF NOT EXISTS applications (
        chat_id INTEGER NOT NULL, job_key TEXT NOT NULL, status TEXT NOT NULL,
        updated REAL NOT NULL, PRIMARY KEY(chat_id, job_key))""")
    db.execute("""CREATE TABLE IF NOT EXISTS resumes (
        chat_id INTEGER PRIMARY KEY, filename TEXT NOT NULL,
        resume_text TEXT NOT NULL, updated REAL NOT NULL)""")
    db.execute("""CREATE TABLE IF NOT EXISTS ai_consent (
        chat_id INTEGER PRIMARY KEY, enabled INTEGER NOT NULL DEFAULT 0)""")


def store_job(db, job):
    db.execute("INSERT OR REPLACE INTO job_snapshots VALUES (?,?,?)",
               (job.key, json.dumps(asdict(job), ensure_ascii=False), time.time()))


def lookup_job(db, prefix):
    if not re.fullmatch(r"[0-9a-f]{16}", prefix or ""):
        return None
    results = db.execute("SELECT payload FROM job_snapshots WHERE job_key LIKE ? LIMIT 2",
                         (prefix + "%",)).fetchall()
    return Job(**json.loads(results[0][0])) if len(results) == 1 else None


def put_feedback(db, chat_id, job, vote):
    if vote not in ALLOWED_FEEDBACK:
        raise ValueError("Unknown feedback")
    store_job(db, job)
    db.execute("INSERT OR REPLACE INTO feedback VALUES (?,?,?,?)",
               (chat_id, job.key, vote, time.time()))


def get_feedback(db, chat_id):
    return {key: vote for key, vote in db.execute(
        "SELECT job_key, vote FROM feedback WHERE chat_id=?", (chat_id,))}


def set_application(db, chat_id, job, status):
    if status not in APPLICATION_STATUSES:
        raise ValueError("Unknown application status")
    store_job(db, job)
    db.execute("INSERT OR REPLACE INTO applications VALUES (?,?,?,?)",
               (chat_id, job.key, status, time.time()))


def get_applications(db, chat_id, limit=20):
    rows = db.execute(
        """SELECT a.job_key, j.payload, a.status FROM applications a
        JOIN job_snapshots j ON j.job_key = a.job_key
        WHERE a.chat_id=? ORDER BY a.updated DESC LIMIT ?""", (chat_id, limit)).fetchall()
    return [(Job(**json.loads(payload)), status) for key, payload, status in rows]


def resume_save(db, chat_id, filename, content):
    if len(content) < 150:
        raise ValueError("Resume needs at least 150 characters")
    if len(content) > 25000:
        raise ValueError("Resume exceeds 25000 characters")
    db.execute("INSERT OR REPLACE INTO resumes VALUES (?,?,?,?)",
               (chat_id, filename[:180], content, time.time()))


def resume_get(db, chat_id):
    row = db.execute("SELECT filename, resume_text FROM resumes WHERE chat_id=?", (chat_id,)).fetchone()
    return tuple(row) if row else None


def resume_remove(db, chat_id):
    db.execute("DELETE FROM resumes WHERE chat_id=?", (chat_id,))


def consent_set(db, chat_id, enabled):
    db.execute("INSERT OR REPLACE INTO ai_consent VALUES (?, ?)", (chat_id, int(enabled)))


def consent_get(db, chat_id):
    row = db.execute("SELECT enabled FROM ai_consent WHERE chat_id=?", (chat_id,)).fetchone()
    return bool(row and row[0])


def topics(job):
    title = ((job.title or "") + " " + (job.company or "")).lower()
    return {name for name, pattern in TOPICS.items() if re.search(pattern, title, re.I)}


def feedback_bonus(job, db, chat_id):
    """At most ±12 points; never overrides a stronger tier in the sort key."""
    signature = topics(job)
    if not signature:
        return 0
    rows = db.execute(
        """SELECT f.vote, j.payload FROM feedback f
        JOIN job_snapshots j ON j.job_key=f.job_key WHERE f.chat_id=?""", (chat_id,)).fetchall()
    value = 0
    for vote, payload in rows:
        old_job = Job(**json.loads(payload))
        overlap = signature.intersection(topics(old_job))
        if overlap:
            value += min(3, len(overlap)) * (2 if vote == "like" else -2)
    return max(-12, min(12, value))


def parse_resume(filename, file_bytes):
    if len(file_bytes) > 5 * 1024 * 1024:
        raise ValueError("Максимальный размер резюме — 5 МБ")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "txt":
        content = file_bytes.decode("utf-8-sig")
    elif ext == "docx":
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        content = "\n".join(paragraphs)
    elif ext == "pdf":
        reader = PdfReader(io.BytesIO(file_bytes), strict=False)
        if len(reader.pages) > 12:
            raise ValueError("Резюме: максимум 12 страниц")
        content = "\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        raise ValueError("Пришли резюме .pdf, .docx или .txt")
    content = "\n".join(line.strip() for line in content.splitlines() if line.strip())
    if len(content) < 150:
        raise ValueError("Не удалось извлечь текст. Для сканированного PDF пришли DOCX или TXT")
    return content[:25000]


def offline_comparison(job, resume):
    from scoring import classify
    fit = classify(job)
    keys = {
        "SQL": r"\bsql\b", "Python": r"\bpython\b", "A/B-тесты": r"a/b.?тест|a/b.?test|ab.?test",
        "когортный анализ": r"когорт|cohort", "конверсия": r"конверс|conversion",
        "Retention": r"retention|удержан|возвращающ", "GMV": r"\bgmv\b",
        "BI": r"дашборд|dashboard|\bbi\b", "финансовые показатели": r"маржин|profit|gmv|прибыл",
    }
    present = [name for name, pattern in keys.items()
               if re.search(pattern, job.description + " " + job.title, re.I)]
    matched = [name for name in present if re.search(keys[name], resume, re.I)]
    unverified = [name for name in present if name not in matched]
    def joining(items):
        return ", ".join(items[:7]) if items else "не удалось определить"
    return (
        f"📊 Предварительное сравнение по тексту (не AI)\n"
        f"{fit.label}, {fit.score}/100 — оценка по вакансии, не вероятность оффера.\n\n"
        f"✅ Упоминаются и в резюме, и в вакансии: {joining(matched)}.\n"
        f"❓ Есть в вакансии, но не подтверждены резюме: {joining(unverified)}.\n\n"
        f"Уточни реальную возможность worldwide remote, язык команды и зарплату net."
    )


def offline_cover(job, resume):
    skills = []
    for label, pattern in (
        ("SQL", r"\bsql\b"), ("Python", r"\bpython\b"),
        ("Excel", r"\bexcel\b"), ("когортным анализом", r"когорт|cohort"),
        ("продуктовыми метриками", r"продуктов|конверс|retention|conversion"),
        ("коммерческими метриками", r"коммерч|\bgmv\b|маржин"),
    ):
        if re.search(pattern, resume, re.I):
            skills.append(label)
    if skills:
        evidence = "В моём резюме отражена работа с " + ", ".join(skills[:4]) + "."
    else:
        evidence = "Мой релевантный опыт описан в приложенном резюме."
    return (
        f"Здравствуйте, команда {job.company or 'компании'}!\n\n"
        f"Меня заинтересовала вакансия «{job.title}». "
        f"Хочу развиваться в аналитике и помогать продуктовым и бизнес-командам "
        f"принимать решения на основе данных. {evidence} "
        f"Буду рад обсудить, как мой опыт может быть полезен вашей команде.\n\n"
        f"Спасибо за рассмотрение отклика!\n\n"
        "Черновик: проверьте формулировки и добавьте личную мотивацию перед отправкой."
    )


def make_docx(title, body):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(9.5)
    document.add_heading("РЕЗЮМЕ — ЧЕРНОВИК", 0)
    document.add_paragraph("Адаптация под: " + title)
    document.add_paragraph("Требуется проверить факты и оформление перед отправкой работодателю.")
    for line in body.splitlines():
        line = line.strip()
        if not line:
            continue
        if len(line) < 80 and (line.isupper() or line.startswith(("## ", "# "))):
            document.add_heading(line.lstrip("# ").strip(), level=2)
        else:
            document.add_paragraph(line, style="Normal")
    file = io.BytesIO()
    document.save(file)
    return file.getvalue()


def make_pdf(title, body):
    path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    pdfmetrics.registerFont(TTFont("ResumeDejaVu", path))
    output = io.BytesIO()
    style = ParagraphStyle("Body", fontName="ResumeDejaVu", fontSize=9,
                           leading=13, spaceAfter=6, wordWrap="CJK")
    heading = ParagraphStyle("Heading", parent=style, fontSize=13, leading=17,
                             spaceAfter=10, textColor=colors.HexColor("#16324F"))
    small = ParagraphStyle("Small", parent=style, fontSize=8, leading=11,
                           textColor=colors.grey, spaceAfter=10)
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=42, rightMargin=42,
                                 topMargin=35, bottomMargin=35)
    story = [Paragraph("РЕЗЮМЕ — ЧЕРНОВИК", heading),
             Paragraph("Вакансия: " + escape(title), style),
             Paragraph("Перед отправкой проверьте соответствие фактов исходному резюме.", small),
             Spacer(1, 4)]
    for raw in body.splitlines():
        line = raw.strip()
        if not line:
            continue
        special = len(line) < 75 and (line.isupper() or line.startswith(("# ", "## ")))
        story.append(Paragraph(escape(line.lstrip("# ").strip()), heading if special else style))
    document.build(story)
    return output.getvalue()
